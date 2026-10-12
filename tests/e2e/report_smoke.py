#!/usr/bin/env python3
"""Actual Compose Solver → parser → PNG → PPTX, versioning and stage recovery."""

import base64
import concurrent.futures
from decimal import Decimal, ROUND_HALF_UP
import hashlib
import io
import json
import os
from pathlib import Path
import re
import subprocess
import sys
import uuid

from PIL import Image
from pptx import Presentation

from agent_smoke import Api, assert_identity
from messaging_smoke import LocalServices
from parser_smoke import artifacts, download, metrics, until
from reconciliation_smoke import queue_message
from scheduler_smoke import licenses


def progress(message):
    print("[M8] " + message, file=sys.stderr, flush=True)


def user(name, password):
    return Api(
        "http://127.0.0.1:8080",
        {
            "Authorization": "Basic "
            + base64.b64encode(f"{name}:{password}".encode()).decode()
        },
    )


def get_job(api, job):
    return api.require("GET", f"/api/v1/jobs/{job['jobNo']}")


def create(api, entry, owned):
    created = api.require(
        "POST",
        "/api/v1/jobs",
        {
            "projectId": 1 if entry["analysisType"] == "STRUCTURE" else 2,
            "modelRevisionId": {"V36": 1, "V37": 2, "P1": 3}[entry["modelRevision"]],
            "analysisType": entry["analysisType"],
            "conditionCode": entry["conditionCode"],
            "casePath": entry["casePath"],
            "resourceProfile": "SMALL",
            "mock": {
                "runtimeSeconds": 2,
                "successProbability": 1,
                "failureMode": "NONE",
                "progressInterval": 0.1,
            },
        },
    )
    job = api.require("GET", f"/api/v1/jobs/{created['jobNo']}")
    owned.append(job)
    return job


def submit(api, job):
    api.require(
        "POST",
        f"/api/v1/jobs/{job['jobNo']}/submit",
        {"idempotencyKey": "m8-" + uuid.uuid4().hex},
    )


def success(api, job, timeout=90):
    def probe():
        current = get_job(api, job)
        if current["status"] in {
            "VALIDATION_FAILED",
            "SUBMIT_FAILED",
            "RUNTIME_FAILED",
            "POST_PROCESS_FAILED",
            "REPORT_FAILED",
            "CANCELLED",
        }:
            raise AssertionError(
                f"{job['jobNo']} failed: {current['failureCode']} / {current['failureMessage']}"
            )
        return current

    return until(
        probe,
        lambda row: row["status"] == "SUCCESS",
        f"{job['jobNo']} complete report",
        timeout,
    )


def report_version(api, job, parser="v1", version=1):
    return until(
        lambda: api.require("GET", f"/api/v1/jobs/{job['jobNo']}/reports"),
        lambda rows: any(
            row["parserVersion"] == parser and row["reportVersion"] == version
            for row in rows
        ),
        f"{job['jobNo']} {parser} report {version}",
        90,
    )


def validate_report(api, job, parser="v1", version=1):
    report = next(
        row
        for row in report_version(api, job, parser, version)
        if row["parserVersion"] == parser and row["reportVersion"] == version
    )
    context = report["context"]
    all_files = artifacts(api, job)
    selected = [
        row
        for row in all_files
        if row["parserVersion"] == parser and row["stageVersion"] == version
    ]
    images = [row for row in selected if row["artifactType"] == "IMAGE"]
    assert len(images) == 4
    hashes = set()
    output = Path(".local/acceptance") / job["jobNo"] / parser / f"report-v{version}"
    output.mkdir(parents=True, exist_ok=True)
    for row in images:
        data = download(api, row)
        with Image.open(io.BytesIO(data)) as image:
            assert image.format == "PNG" and image.size == (1200, 720)
            assert (
                image.info["Description"]
                == f"{job['conditionCode']};subcase={context['subcase']};{row['metadata']['resultType']}"
            )
            image.verify()
        assert row["metadata"]["conditionCode"] == context["conditionCode"]
        assert row["metadata"]["subcase"] == context["subcase"]
        hashes.add(hashlib.sha256(data).hexdigest())
        (output / row["fileName"]).write_bytes(data)
    artifact = next(row for row in selected if row["id"] == report["artifactId"])
    data = download(api, artifact)
    (output / "report.pptx").write_bytes(data)
    presentation = Presentation(io.BytesIO(data))
    assert len(presentation.slides) >= 4
    embedded, text = set(), []
    for slide in presentation.slides:
        content = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                content.append(shape.text)
            if shape.has_table:
                content.extend(
                    cell.text for row in shape.table.rows for cell in row.cells
                )
            if shape.shape_type == 13:
                embedded.add(hashlib.sha256(shape.image.blob).hexdigest())
        assert any("DRAFT / Engineering review required" in value for value in content)
        assert any(
            f"{context['conditionCode']} / Subcase {context['subcase']}" in value
            for value in content
        )
        text.extend(content)
    assert embedded == hashes
    assert context["context"]["modelRevision"] in "\n".join(text)
    canonical = metrics(api, job, parser)
    assert len(canonical) == len(context["metrics"])
    for metric in context["metrics"]:
        assert metric["metricCode"] in text
        assert metric["unit"] in text
        assert (
            metric["valueNum"]
            if metric["valueNum"] is not None
            else metric["valueText"]
        ) in text
        assert (
            metric["conditionCode"] == context["conditionCode"]
            and metric["subcase"] == context["subcase"]
        )
    manifest = next(row for row in selected if row["artifactType"] == "REPORT_MANIFEST")
    metadata = json.loads(download(api, manifest))
    assert (
        set(metadata["imageChecksums"].values()) == hashes
        and metadata["reportStatus"] == "DRAFT"
    )
    return report, {
        "jobNo": job["jobNo"],
        "parserVersion": parser,
        "reportVersion": version,
        "reportId": report["id"],
        "conditionCode": context["conditionCode"],
        "subcase": context["subcase"],
        "slides": len(presentation.slides),
        "canonicalMetrics": len(canonical),
        "pptxSha256": hashlib.sha256(data).hexdigest(),
        "imageSha256": sorted(hashes),
        "localArtifacts": str(output),
    }


def identity(api, workers, job):
    history = api.require("GET", f"/api/v1/jobs/{job['jobNo']}/executions")
    assert len(history) == 1 and history[0]["status"] == "SUCCESS"
    execution = history[0]
    remote = workers[execution["nodeId"]].require(
        "GET", "/worker/v1/jobs/by-request/" + execution["submitRequestId"]
    )
    assert_identity(
        {
            "executionRequestId": execution["submitRequestId"],
            **{
                key: execution[key]
                for key in ("externalExecutionId", "pid", "processIdentity")
            },
        },
        remote,
    )
    return execution, remote


def file_stats(services, node, paths):
    for path in paths:
        assert re.fullmatch(
            r"jobs/CAE-\d{8}-\d{6,}/derived/v[12]/post-v\d+/(contour\.png|mode-shape\.png|comparison\.png|curve\.png|post-manifest\.json)",
            path,
        )
    code = "import sys,json,hashlib;from pathlib import Path;r=Path('/data');names=json.load(sys.stdin);print(json.dumps({n:{'mtime':(r/n).stat().st_mtime_ns,'sha256':hashlib.sha256((r/n).read_bytes()).hexdigest()} for n in names}))"
    result = subprocess.run(
        [*services.compose, "exec", "-T", "worker-" + node, "python", "-c", code],
        input=json.dumps(paths),
        capture_output=True,
        text=True,
        cwd=services.repo,
        env=services.environment,
        timeout=10,
        check=True,
    )
    return json.loads(result.stdout)


def fault(workers, stage, job, count=0, delay=0):
    for api in workers.values():
        api.require(
            "POST",
            "/worker/v1/faults/render",
            {
                "stage": stage,
                "jobNo": job["jobNo"],
                "failuresRemaining": count,
                "delaySeconds": delay,
            },
        )


def verify_comparison(current, baseline):
    before = {row["metricCode"]: row for row in baseline["context"]["metrics"]}
    for row in current["context"]["comparisons"]:
        expected = Decimal(row["current"]) - Decimal(
            before[row["metricCode"]]["valueNum"]
        )
        assert Decimal(row["delta"]) == expected
        percent = (
            expected / Decimal(before[row["metricCode"]]["valueNum"]) * 100
        ).quantize(Decimal("0.000001"), rounding=ROUND_HALF_UP)
        assert Decimal(row["deltaPercent"]) == percent
    assert current["baselineJobNo"] == baseline["jobNo"]


def failure_case(api, services, workers, entry, owned, stage):
    job = create(api, entry, owned)
    fault(workers, stage, job, count=8)
    submit(api, job)
    consumer = "post-process-v1" if stage == "post" else "draft-report-v1"
    rows = until(
        lambda: services.query(
            "SELECT JSON_OBJECT('eventId',logical_event_id,'status',status,'retryCount',retry_count) "
            f"FROM cae_message_inbox WHERE job_id={int(job['id'])} AND consumer_name='{consumer}'"
        ),
        lambda rows: any(row["status"] == "DEAD" for row in rows),
        f"{job['jobNo']} bounded {stage} failure",
        90,
    )
    dead = next(row for row in rows if row["status"] == "DEAD")
    assert dead["retryCount"] == 4
    failed = get_job(api, job)
    assert failed["status"] == (
        "POST_PROCESS_FAILED" if stage == "post" else "REPORT_FAILED"
    )
    execution, original = identity(api, workers, job)
    images = [
        row["uri"] for row in artifacts(api, job) if row["artifactType"] == "IMAGE"
    ]
    old = file_stats(services, execution["nodeId"], images) if images else {}
    fault(workers, stage, job)
    key = "m8-retry-" + uuid.uuid4().hex
    for _ in range(2):
        api.require(
            "POST", f"/api/v1/jobs/{job['jobNo']}/retry", {"idempotencyKey": key}
        )
    success(api, job)
    report, checked = validate_report(api, job)
    _, replay = identity(api, workers, job)
    assert_identity(original, replay)
    assert file_stats(services, execution["nodeId"], images) == old if images else True
    return {
        "jobNo": job["jobNo"],
        "failedStage": stage,
        "logicalEventId": dead["eventId"],
        "boundedAttempts": 4,
        "sameExecutionId": execution["id"],
        "sameProcessIdentity": replay["processIdentity"],
        "postImagesPreserved": stage == "report",
        "reportId": report["id"],
        "validation": checked,
    }


def restart_case(api, services, workers, entry, owned):
    job = create(api, entry, owned)
    fault(workers, "post", job, delay=10)
    submit(api, job)
    until(
        lambda: api.require("GET", f"/api/v1/jobs/{job['jobNo']}/pipeline"),
        lambda rows: any(
            row["kind"] == "POST" and row["status"] == "RUNNING" for row in rows
        ),
        "render running before backend restart",
        60,
    )
    execution, original = identity(api, workers, job)
    prefix = f"jobs/{job['jobNo']}/derived/v1/post-v1/"
    paths = [
        prefix + name
        for name in (
            "contour.png",
            "mode-shape.png",
            "comparison.png",
            "curve.png",
            "post-manifest.json",
        )
    ]
    worker_ids = subprocess.run(
        [
            *services.compose,
            "ps",
            "-q",
            "worker-node-a",
            "worker-node-b",
            "worker-node-c",
        ],
        capture_output=True,
        text=True,
        cwd=services.repo,
        env=services.environment,
        check=True,
    ).stdout
    progress("Restarting backend while the original Worker publishes its post output")
    subprocess.run(
        [*services.compose, "restart", "--no-deps", "backend"],
        capture_output=True,
        text=True,
        cwd=services.repo,
        env=services.environment,
        timeout=45,
        check=True,
    )

    # Completed output can exist before the expired old lease is reclaimed by the backend.
    def published():
        try:
            return file_stats(services, execution["nodeId"], paths)
        except subprocess.CalledProcessError:
            return None

    original_files = until(
        published,
        lambda value: value is not None,
        "original Worker post publication",
        45,
    )
    until(
        lambda: api.call("GET", "/actuator/health")[1],
        lambda body: body.get("status") == "UP",
        "restarted backend health",
        90,
    )
    success(api, job, timeout=150)
    _, checked = validate_report(api, job)
    assert file_stats(services, execution["nodeId"], paths) == original_files
    _, recovered = identity(api, workers, job)
    assert_identity(original, recovered)
    current_workers = subprocess.run(
        [
            *services.compose,
            "ps",
            "-q",
            "worker-node-a",
            "worker-node-b",
            "worker-node-c",
        ],
        capture_output=True,
        text=True,
        cwd=services.repo,
        env=services.environment,
        check=True,
    ).stdout
    assert worker_ids == current_workers
    stages = api.require("GET", f"/api/v1/jobs/{job['jobNo']}/pipeline")
    assert next(row for row in stages if row["kind"] == "POST")["claimGeneration"] >= 2
    fault(workers, "post", job)
    return {
        "jobNo": job["jobNo"],
        "sameExecutionId": execution["id"],
        "sameProcessIdentity": recovered["processIdentity"],
        "workerContainersUnchanged": True,
        "postHashesAndMtimesUnchanged": True,
        "reclaimedStageLease": True,
        "validation": checked,
    }


def main():
    services = LocalServices()
    engineer = user(
        os.getenv("CAE_API_USER", "engineer"),
        os.getenv("CAE_API_PASSWORD", "engineer123"),
    )
    lead = user("lead", "lead123")
    admin = user("admin", "admin123")
    workers = {
        f"node-{name}": Api(
            f"http://127.0.0.1:{port}",
            {"X-Service-Token": os.getenv("SERVICE_TOKEN", "demo-service-token")},
        )
        for name, port in zip("abc", (8001, 8002, 8003), strict=True)
    }
    broker = Api("http://127.0.0.1:15672", services.rabbit_auth())
    baseline_ledger = licenses(engineer)
    assert all(reserved == used == 0 for reserved, used in baseline_ledger.values())
    historical = json.loads(
        (services.repo / "docs/validation/m5-messaging-smoke.json").read_text()
    )["deadLetter"]["logicalEventId"]
    assert queue_message(broker, "cae.job.submit.dlq", historical)
    cases = {row["casePath"]: row for row in engineer.require("GET", "/api/v1/cases")}
    owned, evidence = [], {}
    try:
        progress("Running V36 baseline through real Solver, parser, PNG and PPTX")
        baseline = create(engineer, cases["structure/case-baseline"], owned)
        submit(engineer, baseline)
        success(engineer, baseline)
        baseline_report, evidence["baseline"] = validate_report(engineer, baseline)
        progress("Running V37 current and checking backend comparison arithmetic")
        current = create(engineer, cases["structure/case-current"], owned)
        submit(engineer, current)
        success(engineer, current)
        current_report, evidence["current"] = validate_report(engineer, current)
        verify_comparison(current_report, baseline_report)
        progress("Running actual BIW_MODAL / Subcase 2 report")
        modal = create(engineer, cases["structure/case-modal"], owned)
        submit(engineer, modal)
        success(engineer, modal)
        _, evidence["modal"] = validate_report(engineer, modal)
        assert evidence["modal"]["subcase"] == 2
        execution, original = identity(engineer, workers, current)
        assert (
            engineer.call("POST", f"/api/v1/reports/{current_report['id']}/review", {})[
                0
            ]
            == 403
        )
        assert (
            lead.call("POST", f"/api/v1/reports/{current_report['id']}/approve", {})[0]
            == 403
        )
        for _ in range(2):
            lead.require("POST", f"/api/v1/reports/{current_report['id']}/review", {})
            admin.require("POST", f"/api/v1/reports/{current_report['id']}/approve", {})
        progress("Reparsing v2 without another Solver or changing approved v1 bytes")
        with concurrent.futures.ThreadPoolExecutor(max_workers=6) as pool:
            replies = list(
                pool.map(
                    lambda _: engineer.require(
                        "POST",
                        f"/api/v1/jobs/{current['jobNo']}/reparse",
                        {"parserVersion": "v2"},
                    ),
                    range(6),
                )
            )
        assert len({row["id"] for row in replies}) == 1
        _, evidence["reparseV2"] = validate_report(engineer, current, "v2")
        _, replay = identity(engineer, workers, current)
        assert_identity(original, replay)
        retained = engineer.require("GET", f"/api/v1/reports/{current_report['id']}")
        assert retained["reportStatus"] == "APPROVED"
        _, old = validate_report(engineer, current)
        assert old["pptxSha256"] == evidence["current"]["pptxSha256"]
        progress("Generating another draft version with an explicit baseline")
        request = {
            "parserVersion": "v1",
            "baselineJobNo": baseline["jobNo"],
            "idempotencyKey": "m8-new-draft-" + uuid.uuid4().hex,
        }
        replies = [
            lead.require("POST", f"/api/v1/jobs/{current['jobNo']}/reports", request)
            for _ in range(2)
        ]
        assert replies[0]["id"] == replies[1]["id"]
        new, evidence["newDraft"] = validate_report(engineer, current, "v1", 2)
        assert new["reportStatus"] == "DRAFT"
        verify_comparison(new, baseline_report)
        assert get_job(engineer, current)["status"] == "SUCCESS"
        progress("Injecting owned post and report failures, then stage-only retry")
        evidence["postRetry"] = failure_case(
            engineer, services, workers, cases["structure/case-current"], owned, "post"
        )
        evidence["reportRetry"] = failure_case(
            engineer, services, workers, cases["structure/case-modal"], owned, "report"
        )
        evidence["backendRestart"] = restart_case(
            engineer, services, workers, cases["structure/case-current"], owned
        )
    finally:
        for stage in ("post", "report"):
            for api in workers.values():
                api.require(
                    "POST",
                    "/worker/v1/faults/render",
                    {"stage": stage, "failuresRemaining": 0, "delaySeconds": 0},
                )
    until(
        lambda: licenses(engineer),
        lambda value: value == baseline_ledger,
        "resource baseline restoration",
        30,
    )
    assert queue_message(broker, "cae.job.submit.dlq", historical)
    evidence.update(
        {
            "milestone": 8,
            "historicalDeadLetterPreserved": historical,
            "licenseLedgerReturnedToBaseline": True,
            "realSolverJobCount": len(owned),
            "sameCurrentExecutionId": execution["id"],
        }
    )
    print(json.dumps(evidence, ensure_ascii=False))


if __name__ == "__main__":
    main()
