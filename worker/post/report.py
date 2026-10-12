"""Draft reports display authoritative KPI/comparison values and verified image bytes."""

import io
from datetime import datetime, timedelta, timezone
from typing import Protocol

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


class ReportGenerator(Protocol):
    def generate(self, request, images, path) -> int: ...


class PptxDraftReportGenerator:
    def generate(self, request, images, path):
        presentation = Presentation()
        presentation.slide_width = Inches(13.33)
        presentation.slide_height = Inches(7.5)

        def text(slide, content, x, y, w, h, size=18, color="243247"):
            frame = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
            frame.word_wrap = True
            for index, line in enumerate(content.split("\n")):
                paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
                paragraph.text = line
                paragraph.font.size = Pt(size)
                paragraph.font.color.rgb = RGBColor.from_string(color)

        def slide(title):
            current = presentation.slides.add_slide(presentation.slide_layouts[6])
            text(current, title, 0.5, 0.25, 12.3, 0.7, 28)
            text(
                current,
                f"{request.condition_code} / Subcase {request.subcase} / {request.parser_version}",
                0.5,
                1,
                12.3,
                0.4,
                14,
            )
            text(
                current,
                "DRAFT / Engineering review required · SYNTHETIC PUBLIC DEMO",
                0.5,
                7.05,
                12.3,
                0.3,
                12,
                "9A3412",
            )
            return current

        def table(current, headers, rows, x=0.5, y=1.7, width=12.3, height=4.8):
            shape = current.shapes.add_table(
                len(rows) + 1, len(headers), Inches(x), Inches(y), Inches(width), Inches(height)
            )
            for row_index, row in enumerate([headers, *rows]):
                for column, value in enumerate(row):
                    cell = shape.table.cell(row_index, column)
                    cell.text = str(value) if value is not None else "N/A"
                    for paragraph in cell.text_frame.paragraphs:
                        paragraph.font.size = Pt(13)

        current = slide(
            "CAE simulation result · Draft report"
            if request.template_code == "STRUCTURE_DRAFT"
            else "CAE pilot summary · Draft report"
        )
        context = request.context
        generated = datetime.fromisoformat(request.generated_at.replace("Z", "+00:00"))
        generated = generated.astimezone(timezone(timedelta(hours=8)))
        text(
            current,
            f"Project: {context.project_code} / {context.project_name}\n"
            f"Model: {context.model_code} / {context.model_name}\n"
            f"Revision: {context.model_revision}\n"
            f"Analysis: {request.analysis_type} · Maturity: {context.maturity}\n"
            f"Job: {request.job_no}\n"
            f"Execution: {request.execution_id} / Parse: {request.parse_record_id}\n"
            f"Report version: {request.report_version} · "
            f"Template: {request.template_code}/{request.template_version}\n"
            f"Generated: {generated:%Y-%m-%d %H:%M:%S %z} (Asia/Shanghai)",
            0.8,
            1.8,
            11.5,
            4.7,
            21,
        )
        current = slide("Canonical KPI summary")
        table(
            current,
            ["Metric", "Value", "Unit", "Status"],
            [
                [
                    m.metric_code,
                    m.value_num if m.value_num is not None else m.value_text,
                    m.unit,
                    m.status,
                ]
                for m in request.metrics
            ],
        )
        current = slide("Historical baseline comparison")
        text(
            current,
            "Baseline: " + (request.baseline.job_no if request.baseline else "N/A"),
            0.5,
            1.45,
            12,
            0.3,
            13,
        )
        table(
            current,
            ["Metric", "Baseline", "Current", "Delta", "Delta %"],
            [
                [m.metric_code, m.baseline, m.current, m.delta, m.delta_percent]
                for m in request.comparisons
            ],
            width=8,
            height=4.7,
            y=1.9,
        )
        current.shapes.add_picture(
            io.BytesIO(images["comparison.png"]), Inches(8.7), Inches(2), width=Inches(4.1)
        )
        current = slide("Synthetic contour and mode shape")
        current.shapes.add_picture(
            io.BytesIO(images["contour.png"]), Inches(0.5), Inches(1.65), width=Inches(6)
        )
        current.shapes.add_picture(
            io.BytesIO(images["mode-shape.png"]), Inches(6.8), Inches(1.65), width=Inches(6)
        )
        current = slide("Synthetic response curve · Conclusion draft")
        current.shapes.add_picture(
            io.BytesIO(images["curve.png"]), Inches(0.5), Inches(1.65), width=Inches(7.8)
        )
        text(
            current,
            "Conclusion draft\n\nEngineering interpretation and acceptance "
            "remain for the reviewer.\n\nKPI values, image condition and Subcase "
            "refer to the same selected parser version.",
            8.6,
            2,
            4.1,
            3.8,
            18,
        )
        presentation.core_properties.title = (
            f"{request.job_no} / {request.condition_code} / Subcase {request.subcase} / DRAFT"
        )
        presentation.save(path)
        return len(presentation.slides)
