"""Open actual generated files and exercise immutable publication across service restarts."""

import copy
import hashlib
import json
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import pytest
from conftest import prepare_request, wait_terminal
from PIL import Image
from pptx import Presentation
from pydantic import ValidationError

from agents.files import aggregate_checksum, atomic_json, sha256_file
from app.errors import AgentError
from post.contracts import PostRequest, ReportRequest
from post.service import RenderService


def source(manager, settings, modal=False):
    submit = prepare_request(settings.data_root)
    if modal:
        submit.condition_code = "BIW_MODAL"
        path = Path(submit.workspace) / "input" / "manifest.json"
        manifest = json.loads(path.read_text())
        manifest["conditionCode"] = "BIW_MODAL"
        manifest["parameters"]["subcase"] = 2
        atomic_json(path, manifest)
        files = {p.name: sha256_file(p) for p in path.parent.iterdir()}
        submit.snapshot_checksum = aggregate_checksum(files)
        atomic_json(
            path.parent.parent / "input-checksums.json",
            {"files": files, "checksum": submit.snapshot_checksum},
        )
    execution, _ = manager.submit(submit)
    terminal = wait_terminal(manager, execution["externalExecutionId"])
    raw_path = Path(terminal["workspace"]) / "raw_result.json"
    raw = json.loads(raw_path.read_text())
    pairs = [
        ("MASS", raw["mass_kg"], "kg", "MASS"),
        (
            "TORSIONAL_STIFFNESS",
            raw["torsional_stiffness_nm_per_deg"],
            "N*m/deg",
            "TORSIONAL_STIFFNESS",
        ),
        ("BENDING_STIFFNESS", raw["bending_stiffness_n_per_mm"], "N/mm", "BENDING_STIFFNESS"),
    ]
    pairs.extend(
        (f"MODE_{i + 1}_FREQUENCY", value, "Hz", "FREQUENCY")
        for i, value in enumerate(raw["modes_hz"])
    )
    metrics = [
        {
            "metricCode": code,
            "metricName": code,
            "valueNum": str(value),
            "valueText": None,
            "unit": unit,
            "dimension": dimension,
            "status": "VALID",
            "conditionCode": submit.condition_code,
            "subcase": raw["subcase"],
        }
        for code, value, unit, dimension in pairs
    ]
    result = {
        "jobNo": submit.job_no,
        "executionId": 1,
        "executionRequestId": submit.execution_request_id,
        "analysisType": "STRUCTURE",
        "conditionCode": submit.condition_code,
        "subcase": raw["subcase"],
        "rawChecksum": sha256_file(raw_path),
        "parserVersion": "v1",
        "maturity": "COMPLETE",
        "metrics": [{**row, "valueNum": float(row["valueNum"])} for row in metrics],
    }
    output = Path(submit.workspace) / "derived" / "v1"
    output.mkdir(parents=True)
    atomic_json(output / "result.json", result)
    comparisons = [
        {
            "metricCode": row["metricCode"],
            "baseline": None,
            "current": row["valueNum"],
            "delta": None,
            "deltaPercent": None,
            "unit": row["unit"],
            "status": "NO_BASELINE",
        }
        for row in metrics
    ]
    return {
        "schemaVersion": 1,
        "jobNo": submit.job_no,
        "executionId": 1,
        "executionRequestId": submit.execution_request_id,
        "externalExecutionId": terminal["externalExecutionId"],
        "nodeId": settings.node_id,
        "pid": terminal["pid"],
        "processIdentity": terminal["processIdentity"],
        "parseRecordId": 1,
        "parserVersion": "v1",
        "postVersion": 1,
        "analysisType": "STRUCTURE",
        "conditionCode": submit.condition_code,
        "subcase": raw["subcase"],
        "rawChecksum": sha256_file(raw_path),
        "snapshotChecksum": submit.snapshot_checksum,
        "resultJsonChecksum": sha256_file(output / "result.json"),
        "context": {
            "projectCode": "PUBLIC",
            "projectName": "Public demo",
            "modelCode": "BIW",
            "modelName": "Synthetic Body",
            "modelRevision": "V37",
            "maturity": "COMPLETE",
        },
        "metrics": metrics,
        "comparisons": comparisons,
        "baseline": None,
        "specs": [
            {
                "conditionCode": submit.condition_code,
                "resultType": kind,
                "component": "SYNTHETIC_BODY",
                "subcase": raw["subcase"],
                "frame": 0,
                "view": "ISOMETRIC",
                "selectionSet": "PUBLIC_DEMO",
                "outputName": name,
            }
            for kind, name in [
                ("CONTOUR", "contour.png"),
                ("MODE_SHAPE", "mode-shape.png"),
                ("COMPARISON", "comparison.png"),
                ("CURVE", "curve.png"),
            ]
        ],
    }


def report_request(post, response, version=1):
    return ReportRequest.model_validate(
        {
            key: value
            for key, value in {
                **post,
                "reportVersion": version,
                "templateCode": "STRUCTURE_DRAFT",
                "templateVersion": "v1",
                "generatedAt": "2026-10-12T03:00:00Z",
                "postManifestChecksum": response["manifest"]["checksum"],
                "postRequestFingerprint": response["requestFingerprint"],
            }.items()
            if key != "specs"
        }
    )


@pytest.mark.parametrize("modal", [False, True])
def test_real_png_and_pptx_values_images_conditions_and_restart_replay(manager, settings, modal):
    payload = source(manager, settings, modal)
    service = RenderService(settings)
    request = PostRequest.model_validate(payload)
    post = service.render(request, "post")
    assert len(post["artifacts"]) == 4
    hashes = set()
    timestamps = {}
    for item in post["artifacts"]:
        path = settings.data_root / item["relativePath"]
        with Image.open(path) as image:
            assert image.size == (1200, 720)
            assert image.getextrema()[0][0] != image.getextrema()[0][1]
        hashes.add(sha256_file(path))
        timestamps[item["relativePath"]] = path.stat().st_mtime_ns
        assert item["metadata"]["subcase"] == payload["subcase"]
    assert len(hashes) == 4
    assert RenderService(settings).render(request, "post") == post
    assert all(
        (settings.data_root / name).stat().st_mtime_ns == time for name, time in timestamps.items()
    )
    report = service.render(report_request(payload, post), "report")
    path = settings.data_root / report["artifacts"][0]["relativePath"]
    presentation = Presentation(path)
    assert len(presentation.slides) == 5
    embedded = set()
    text = []
    for slide in presentation.slides:
        content = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                content.append(shape.text)
            if shape.has_table:
                content.extend(cell.text for row in shape.table.rows for cell in row.cells)
            if shape.shape_type == 13:
                embedded.add(hashlib.sha256(shape.image.blob).hexdigest())
        assert any("DRAFT / Engineering review required" in value for value in content)
        assert any(
            f"{payload['conditionCode']} / Subcase {payload['subcase']}" in value
            for value in content
        )
        text.extend(content)
    assert embedded == hashes
    for metric in payload["metrics"]:
        assert metric["metricCode"] in text and metric["valueNum"] in text
    before = path.stat().st_mtime_ns
    assert RenderService(settings).render(report_request(payload, post), "report") == report
    assert path.stat().st_mtime_ns == before
    assert len(manager.registry.all()) == 1


def test_concurrent_duplicates_publish_once_and_changed_fingerprint_cannot_overwrite(
    manager, settings
):
    payload = source(manager, settings)
    request = PostRequest.model_validate(payload)
    service = RenderService(settings)
    with ThreadPoolExecutor(max_workers=4) as pool:
        replies = list(pool.map(lambda _: service.render(request, "post"), range(8)))
    assert all(row == replies[0] for row in replies)
    changed = copy.deepcopy(payload)
    changed["context"]["projectName"] = "Different frozen context"
    with pytest.raises(AgentError, match="different frozen inputs"):
        service.render(PostRequest.model_validate(changed), "post")
    path = settings.data_root / replies[0]["artifacts"][0]["relativePath"]
    path.write_bytes(b"corrupt")
    with pytest.raises(AgentError, match="checksum mismatch"):
        service.render(request, "post")
    assert path.read_bytes() == b"corrupt"


@pytest.mark.parametrize(
    "field", ["subcase", "condition", "pid", "checksum", "metric", "symlink", "unknown"]
)
def test_invalid_source_or_spec_never_publishes(manager, settings, field):
    payload = source(manager, settings)
    if field == "subcase":
        payload["specs"][0]["subcase"] = 2
    elif field == "condition":
        payload["conditionCode"] = "OTHER"
    elif field == "pid":
        payload["pid"] += 1
    elif field == "checksum":
        payload["rawChecksum"] = "f" * 64
    elif field == "metric":
        payload["metrics"][0]["valueNum"] = "999"
        payload["comparisons"][0]["current"] = "999"
    elif field == "unknown":
        payload["workspace"] = "/tmp/uncontrolled"
    else:
        path = settings.data_root / "jobs" / payload["jobNo"] / "derived" / "v1" / "result.json"
        target = settings.data_root / "outside.json"
        target.write_bytes(path.read_bytes())
        path.unlink()
        path.symlink_to(target)
    with pytest.raises((AgentError, ValidationError)):
        RenderService(settings).render(PostRequest.model_validate(payload), "post")
    assert not (
        settings.data_root / "jobs" / payload["jobNo"] / "derived" / "v1" / "post-v1"
    ).exists()


def test_failed_render_cleans_private_output_and_recovers_same_source(manager, settings):
    payload = source(manager, settings)
    service = RenderService(settings)
    service.fault("post", 1)
    with pytest.raises(AgentError, match="Injected"):
        service.render(PostRequest.model_validate(payload), "post")
    post = service.render(PostRequest.model_validate(payload), "post")
    service.fault("report", 1)
    with pytest.raises(AgentError, match="Injected"):
        service.render(report_request(payload, post), "report")
    assert service.render(report_request(payload, post), "report")["reportStatus"] == "DRAFT"
    assert len(manager.registry.all()) == 1


def test_report_rejects_different_post_kpi_or_manifest_before_publication(manager, settings):
    payload = source(manager, settings)
    service = RenderService(settings)
    post = service.render(PostRequest.model_validate(payload), "post")
    report = report_request(payload, post)
    report.context.model_revision = "V99"
    with pytest.raises(AgentError, match="selected post context"):
        service.render(report, "report")
    assert not (
        settings.data_root / "jobs" / payload["jobNo"] / "reports" / "v1" / "report-v1"
    ).exists()


def test_service_token_unknown_fields_and_health_during_render(manager, settings):
    import threading

    from fastapi.testclient import TestClient

    from app.main import create_app

    payload = source(manager, settings)
    with TestClient(create_app(settings)) as client:
        assert client.post("/worker/v1/post/process", json=payload).status_code == 401
        headers = {"X-Service-Token": settings.service_token}
        assert (
            client.post(
                "/worker/v1/post/process", headers=headers, json={**payload, "outputPath": "/tmp"}
            ).status_code
            == 422
        )
        assert (
            client.post(
                "/worker/v1/faults/render",
                headers=headers,
                json={"stage": "post", "failuresRemaining": 101},
            ).status_code
            == 422
        )
        started = threading.Event()
        proceed = threading.Event()
        original = client.app.state.render_service.processor.process

        def paused(*arguments):
            started.set()
            assert proceed.wait(5)
            return original(*arguments)

        client.app.state.render_service.processor.process = paused
        with ThreadPoolExecutor(max_workers=1) as pool:
            response = pool.submit(
                client.post, "/worker/v1/post/process", headers=headers, json=payload
            )
            assert started.wait(5)
            try:
                assert client.get("/health").status_code == 200
                assert client.get("/worker/v1/node/status", headers=headers).status_code == 200
            finally:
                proceed.set()
            assert response.result().status_code == 200


def test_canonical_unicode_fingerprint_matches_shared_java_fixture():
    from post.service import fingerprint

    fixture = json.loads(
        (Path(__file__).resolve().parents[2] / "docs/contracts/render-fingerprint.json").read_text()
    )
    assert fingerprint(fixture["payload"]) == fixture["sha256"]
